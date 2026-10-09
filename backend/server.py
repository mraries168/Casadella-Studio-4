from fastapi import FastAPI, APIRouter, Depends, HTTPException, Query, UploadFile, File, Request
from fastapi.responses import Response, HTMLResponse, FileResponse
from starlette.concurrency import run_in_threadpool
import asyncio
import requests
import uuid
import re
import io
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field
from typing import Optional, List, Literal
from datetime import datetime, timezone, timedelta
from pathlib import Path
from bson import ObjectId
import bcrypt
import jwt
import os
import logging

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

client = AsyncIOMotorClient(os.environ['MONGO_URL'])
db = client[os.environ['DB_NAME']]
LOCAL_TZ = timedelta(hours=8)  # WITA (Kupang)
JWT_SECRET = os.environ['JWT_SECRET']
ALGO = "HS256"
TOKEN_HOURS = 12

app = FastAPI()
api = APIRouter(prefix="/api")
bearer = HTTPBearer(auto_error=False)
logger = logging.getLogger(__name__)

ROLE_MAP = {"super_admin": "super_admin", "admin1": "admin", "admin2": "admin",
            "photographer": "staff", "editor": "staff"}


# ---------- helpers ----------
def ser(doc):
    if not doc:
        return doc
    doc = dict(doc)
    doc["id"] = str(doc.pop("_id"))
    if "pin_hash" in doc:
        doc["permissions"] = eff_perms(doc)
    doc.pop("pin_hash", None)
    return doc


def oid(v: str) -> ObjectId:
    try:
        return ObjectId(v)
    except Exception:
        raise HTTPException(404, "Data tidak ditemukan")


def now_iso():
    return datetime.now(timezone.utc).isoformat()


async def audit(actor: str, entity: str, action: str, entity_id: str, summary: str, changes: Optional[list] = None):
    await db.audit_logs.insert_one({"actor": actor, "entity": entity, "action": action, "entity_id": entity_id,
                                    "summary": summary, "changes": changes or [], "created_at": now_iso()})


def diff(old: dict, new: dict, fields: dict) -> list:
    return [{"field": label, "from": old.get(k), "to": new.get(k)} for k, label in fields.items() if old.get(k) != new.get(k)]


def hash_pin(pin: str) -> str:
    return bcrypt.hashpw(pin.encode(), bcrypt.gensalt(rounds=10)).decode()


def to_min(t: str) -> int:
    h, m = t.split(":")
    return int(h) * 60 + int(m)


def to_hhmm(m: int) -> str:
    return f"{m // 60:02d}:{m % 60:02d}"


SETTINGS_DEFAULTS = {"studio_name": "Casadella Photo Studio", "open_time": "09:00", "close_time": "20:00",
                     "slot_interval": 15, "default_dp_percent": 50, "work_start": "09:00", "work_end": "18:00",
                     "late_tolerance_min": 10, "overtime_min_minutes": 30, "overtime_rate_per_hour": 20000,
                     "meal_allowance_per_day": 25000, "rent_auto_enabled": True, "rent_amount": 2500000,
                     "geofence_enabled": True, "studio_lat": -10.158988, "studio_lng": 123.605584, "attendance_radius_m": 100,
                     "admin_wa_number": "6287786098669"}


async def get_settings():
    s = await db.settings.find_one({"key": "system"})
    return {**SETTINGS_DEFAULTS, **ser(s)}


PERMS = ["booking", "expense", "reports"]


def eff_perms(u):
    if u.get("role") == "super_admin":
        return {k: True for k in PERMS}
    default = u.get("role") == "admin"
    p = u.get("permissions") or {}
    return {k: bool(p.get(k, default)) for k in PERMS}


# ---------- auth ----------
async def current_user(request: Request, creds: Optional[HTTPAuthorizationCredentials] = Depends(bearer)):
    token = creds.credentials if creds else request.query_params.get("token")
    if not token:
        raise HTTPException(401, "Tidak terautentikasi")
    try:
        p = jwt.decode(token, JWT_SECRET, algorithms=[ALGO])
        u = await db.users.find_one({"_id": ObjectId(p["sub"])})
    except Exception:
        raise HTTPException(401, "Sesi tidak valid")
    if not u or not u.get("enabled", True):
        raise HTTPException(401, "Pengguna nonaktif")
    return u


def require(*roles):
    async def dep(u=Depends(current_user)):
        if u["role"] not in roles:
            raise HTTPException(403, "Akses ditolak")
        return u
    return dep


ADMIN = require("admin", "super_admin")
SUPER = require("super_admin")


def need(perm):
    async def dep(u=Depends(current_user)):
        if not eff_perms(u)[perm]:
            raise HTTPException(403, "Akses ditolak")
        return u
    return dep


BOOK = need("booking")
EXP = need("expense")
REP = need("reports")


class LoginIn(BaseModel):
    user_id: str
    pin: str = Field(min_length=4, max_length=8)


@api.get("/auth/users")
async def login_users():
    docs = await db.users.find({"enabled": True}, {"pin_hash": 0}).sort("name", 1).to_list(200)
    return [{"id": str(d["_id"]), "name": d["name"], "sub_role": d["sub_role"]} for d in docs]


@api.post("/auth/login")
async def login(body: LoginIn):
    u = await db.users.find_one({"_id": oid(body.user_id)})
    if not u or not u.get("enabled", True) or not bcrypt.checkpw(body.pin.encode(), u["pin_hash"].encode()):
        raise HTTPException(401, "PIN salah")
    exp = datetime.now(timezone.utc) + timedelta(hours=TOKEN_HOURS)
    token = jwt.encode({"sub": str(u["_id"]), "role": u["role"], "exp": exp}, JWT_SECRET, algorithm=ALGO)
    return {"access_token": token, "user": ser(u)}


@api.get("/auth/me")
async def me(u=Depends(current_user)):
    return ser(u)


# ---------- users (super admin) ----------
class UserIn(BaseModel):
    name: str = Field(min_length=1)
    sub_role: Literal["super_admin", "admin1", "admin2", "photographer", "editor"]
    pin: Optional[str] = Field(default=None, min_length=4, max_length=8)
    enabled: bool = True
    phone: str = ""


@api.get("/users")
async def list_users(u=Depends(current_user), sub_role: Optional[str] = None):
    q = {"sub_role": sub_role} if sub_role else {}
    return [ser(d) for d in await db.users.find(q).sort("name", 1).to_list(500)]


@api.post("/users")
async def create_user(body: UserIn, _=Depends(SUPER)):
    if not body.pin or not body.pin.isdigit():
        raise HTTPException(400, "PIN harus berupa angka")
    doc = {"name": body.name.strip(), "sub_role": body.sub_role, "role": ROLE_MAP[body.sub_role],
           "enabled": body.enabled, "phone": body.phone, "pin_hash": hash_pin(body.pin), "created_at": now_iso()}
    r = await db.users.insert_one(doc)
    doc["_id"] = r.inserted_id
    return ser(doc)


@api.put("/users/{uid}")
async def update_user(uid: str, body: UserIn, me_=Depends(SUPER)):
    upd = {"name": body.name.strip(), "sub_role": body.sub_role, "role": ROLE_MAP[body.sub_role], "enabled": body.enabled, "phone": body.phone}
    if body.pin:
        if not body.pin.isdigit():
            raise HTTPException(400, "PIN harus berupa angka")
        upd["pin_hash"] = hash_pin(body.pin)
    if str(me_["_id"]) == uid and body.sub_role != "super_admin":
        raise HTTPException(400, "Tidak dapat menurunkan akses diri sendiri")
    await db.users.update_one({"_id": oid(uid)}, {"$set": upd})
    return ser(await db.users.find_one({"_id": oid(uid)}))


class PermsIn(BaseModel):
    booking: bool
    expense: bool
    reports: bool


@api.put("/users/{uid}/permissions")
async def update_perms(uid: str, body: PermsIn, _=Depends(SUPER)):
    await db.users.update_one({"_id": oid(uid)}, {"$set": {"permissions": body.model_dump()}})
    return ser(await db.users.find_one({"_id": oid(uid)}))


@api.delete("/users/{uid}")
async def delete_user(uid: str, me_=Depends(SUPER)):
    if str(me_["_id"]) == uid:
        raise HTTPException(400, "Tidak dapat menghapus diri sendiri")
    await db.users.delete_one({"_id": oid(uid)})
    return {"ok": True}


# ---------- settings ----------
class SettingsIn(BaseModel):
    studio_name: str
    open_time: str = "09:00"
    close_time: str = "20:00"
    slot_interval: int = 15
    default_dp_percent: int = 50
    work_start: str = "09:00"
    work_end: str = "18:00"
    late_tolerance_min: int = 10
    overtime_min_minutes: int = 30
    overtime_rate_per_hour: float = 20000
    meal_allowance_per_day: float = 25000
    rent_auto_enabled: bool = True
    rent_amount: float = Field(ge=0, default=2500000)
    geofence_enabled: bool = True
    studio_lat: float = Field(ge=-90, le=90, default=-10.158988)
    studio_lng: float = Field(ge=-180, le=180, default=123.605584)
    attendance_radius_m: int = Field(ge=10, le=5000, default=100)
    admin_wa_number: str = Field(default="6287786098669", max_length=20)


@api.get("/settings")
async def read_settings(_=Depends(current_user)):
    return await get_settings()


@api.put("/settings")
async def write_settings(body: SettingsIn, _=Depends(SUPER)):
    await db.settings.update_one({"key": "system"}, {"$set": body.model_dump()}, upsert=True)
    return await get_settings()


# ---------- packages ----------
class PackageIn(BaseModel):
    name: str = Field(min_length=1)
    description: str = ""
    duration_min: int = Field(ge=15)
    price: float = Field(ge=0)
    active: bool = True


@api.get("/packages")
async def list_packages(_=Depends(current_user)):
    return [ser(d) for d in await db.packages.find().sort("price", 1).to_list(200)]


PKG_FIELDS = {"name": "Nama", "description": "Deskripsi", "duration_min": "Durasi (menit)", "price": "Harga", "active": "Aktif"}


@api.post("/packages")
async def create_package(body: PackageIn, u=Depends(ADMIN)):
    doc = body.model_dump()
    r = await db.packages.insert_one(doc)
    doc["_id"] = r.inserted_id
    await audit(u["name"], "package", "create", str(r.inserted_id), f"Paket baru {doc['name']} ({fmt_rp(doc['price'])})")
    return ser(doc)


@api.put("/packages/{pid}")
async def update_package(pid: str, body: PackageIn, u=Depends(ADMIN)):
    old = await db.packages.find_one({"_id": oid(pid)})
    if not old:
        raise HTTPException(404, "Paket tidak ditemukan")
    new = body.model_dump()
    await db.packages.update_one({"_id": oid(pid)}, {"$set": new})
    ch = diff(old, new, PKG_FIELDS)
    if ch:
        await audit(u["name"], "package", "update", pid, f"Ubah paket {new['name']}", ch)
    return ser(await db.packages.find_one({"_id": oid(pid)}))


@api.delete("/packages/{pid}")
async def delete_package(pid: str, u=Depends(ADMIN)):
    old = await db.packages.find_one({"_id": oid(pid)})
    await db.packages.delete_one({"_id": oid(pid)})
    if old:
        await audit(u["name"], "package", "delete", pid, f"Hapus paket {old['name']} ({fmt_rp(old['price'])})")
    return {"ok": True}


# ---------- slots ----------
async def compute_slots(date: str, duration: int, exclude: Optional[str] = None):
    s = await get_settings()
    open_m, close_m, step = to_min(s["open_time"]), to_min(s["close_time"]), s["slot_interval"]
    q = {"date": date, "status": {"$ne": "cancelled"}}
    if exclude:
        q["_id"] = {"$ne": oid(exclude)}
    busy = [(to_min(b["start_time"]), to_min(b["end_time"])) for b in await db.bookings.find(q).to_list(500)]
    out = []
    t = open_m
    while t < close_m:
        end = t + duration
        ok = end <= close_m and all(not (t < be and end > bs) for bs, be in busy)
        out.append({"time": to_hhmm(t), "end": to_hhmm(end), "available": ok})
        t += step
    return out


@api.get("/slots")
async def slots(date: str, package_id: str, exclude_booking: Optional[str] = None, _=Depends(current_user)):
    p = await db.packages.find_one({"_id": oid(package_id)})
    if not p:
        raise HTTPException(404, "Paket tidak ditemukan")
    return await compute_slots(date, p["duration_min"], exclude_booking)


# ---------- bookings ----------
class BookingIn(BaseModel):
    client_name: str = Field(min_length=1)
    client_phone: str = ""
    date: str
    start_time: str
    package_id: str
    photographer_id: Optional[str] = None
    editor_id: Optional[str] = None
    notes: str = ""


class PaymentIn(BaseModel):
    amount: float = Field(gt=0)
    method: Literal["cash", "transfer", "card", "qris"]
    type: Literal["dp", "pelunasan"]
    note: str = ""


def pay_status(paid: float, price: float) -> str:
    if paid <= 0:
        return "belum_bayar"
    return "lunas" if paid >= price else "dp"


async def name_of(uid: Optional[str]):
    if not uid:
        return None
    u = await db.users.find_one({"_id": oid(uid)})
    return u["name"] if u else None


async def build_booking(body: BookingIn, exclude: Optional[str] = None):
    p = await db.packages.find_one({"_id": oid(body.package_id)})
    if not p:
        raise HTTPException(404, "Paket tidak ditemukan")
    sl = await compute_slots(body.date, p["duration_min"], exclude)
    match = next((x for x in sl if x["time"] == body.start_time), None)
    if not match or not match["available"]:
        raise HTTPException(409, "Slot waktu tidak tersedia")
    return {"client_name": body.client_name.strip(), "client_phone": body.client_phone, "date": body.date,
            "start_time": body.start_time, "end_time": match["end"], "package_id": body.package_id,
            "package_name": p["name"], "duration_min": p["duration_min"], "price": p["price"],
            "photographer_id": body.photographer_id, "photographer_name": await name_of(body.photographer_id),
            "editor_id": body.editor_id, "editor_name": await name_of(body.editor_id), "notes": body.notes}


@api.get("/bookings")
async def list_bookings(date: Optional[str] = None, status: Optional[str] = None, _=Depends(current_user)):
    q = {"date": date} if date else {}
    if status:
        q["status"] = status
    return [wf_view(d) for d in await db.bookings.find(q).sort([("date", 1), ("start_time", 1)]).to_list(1000)]


@api.get("/bookings/unpaid-finished")
async def unpaid_finished(_=Depends(current_user)):
    n = datetime.now(timezone.utc) + LOCAL_TZ
    d, t = n.strftime("%Y-%m-%d"), n.strftime("%H:%M")
    q = {"status": "confirmed", "payment_status": {"$ne": "lunas"},
         "$or": [{"date": {"$lt": d}}, {"date": d, "end_time": {"$lte": t}}]}
    return [wf_view(b) for b in await db.bookings.find(q).sort([("date", -1), ("start_time", -1)]).to_list(200)]


@api.get("/bookings/{bid}")
async def get_booking(bid: str, _=Depends(current_user)):
    b = await db.bookings.find_one({"_id": oid(bid)})
    if not b:
        raise HTTPException(404, "Booking tidak ditemukan")
    return wf_view(b)


@api.post("/bookings")
async def create_booking(body: BookingIn, u=Depends(BOOK)):
    doc = await build_booking(body)
    doc.update({"payments": [], "paid_total": 0, "payment_status": "belum_bayar", "status": "confirmed",
                "stage": "order_received", "stage_history": [stage_entry("order_received", u["name"])],
                "created_by": u["name"], "created_at": now_iso()})
    r = await db.bookings.insert_one(doc)
    doc["_id"] = r.inserted_id
    await audit(u["name"], "booking", "create", str(r.inserted_id),
                f"Booking baru {doc['client_name']} · {doc['date']} {doc['start_time']} · {doc['package_name']}")
    await count_auto_task(u, "booking")
    return ser(doc)


BOOKING_FIELDS = {"client_name": "Nama Klien", "client_phone": "No. WhatsApp", "date": "Tanggal", "start_time": "Jam Mulai",
                  "package_name": "Paket", "price": "Harga", "photographer_name": "Fotografer", "editor_name": "Editor", "notes": "Catatan"}


@api.put("/bookings/{bid}")
async def update_booking(bid: str, body: BookingIn, u=Depends(BOOK)):
    old = await db.bookings.find_one({"_id": oid(bid)})
    if not old:
        raise HTTPException(404, "Booking tidak ditemukan")
    doc = await build_booking(body, bid)
    doc["payment_status"] = pay_status(old.get("paid_total", 0), doc["price"])
    await db.bookings.update_one({"_id": oid(bid)}, {"$set": doc})
    ch = diff(old, doc, BOOKING_FIELDS)
    if ch:
        await audit(u["name"], "booking", "update", bid, f"Ubah booking {doc['client_name']} · {doc['date']}", ch)
    return ser(await db.bookings.find_one({"_id": oid(bid)}))


@api.post("/bookings/{bid}/cancel")
async def cancel_booking(bid: str, u=Depends(BOOK)):
    old = await db.bookings.find_one({"_id": oid(bid)})
    r = await db.bookings.delete_one({"_id": oid(bid)})
    if not r.deleted_count:
        raise HTTPException(404, "Booking tidak ditemukan")
    tx = await db.transactions.delete_many({"booking_id": bid})
    await audit(u["name"], "booking", "cancel", bid,
                f"Batalkan booking {old['client_name']} · {old['date']} {old['start_time']} (pembayaran {fmt_rp(old.get('paid_total', 0))} dihapus)")
    return {"ok": True, "deleted_transactions": tx.deleted_count}


@api.post("/bookings/{bid}/approve")
async def approve_booking(bid: str, u=Depends(BOOK)):
    b = await db.bookings.find_one({"_id": oid(bid)})
    if not b or b.get("status") != "pending":
        raise HTTPException(404, "Booking menunggu konfirmasi tidak ditemukan")
    await db.bookings.update_one({"_id": b["_id"]}, {"$set": {"status": "confirmed", "approved_by": u["name"], "approved_at": now_iso(),
                                                              "stage": "order_received"},
                                                     "$push": {"stage_history": stage_entry("order_received", u["name"])}})
    await audit(u["name"], "booking", "approve", bid, f"Setujui booking online {b['client_name']} · {b['date']} {b['start_time']}")
    return ser(await db.bookings.find_one({"_id": b["_id"]}))


@api.post("/bookings/{bid}/balance-requested")
async def mark_balance_requested(bid: str, u=Depends(BOOK)):
    upd = {"balance_requested_at": now_iso(), "balance_requested_by": u["name"]}
    r = await db.bookings.update_one({"_id": oid(bid)}, {"$set": upd})
    if not r.matched_count:
        raise HTTPException(404, "Booking tidak ditemukan")
    return upd


@api.post("/bookings/{bid}/reject")
async def reject_booking(bid: str, u=Depends(BOOK)):
    b = await db.bookings.find_one({"_id": oid(bid)})
    if not b or b.get("status") != "pending":
        raise HTTPException(404, "Booking menunggu konfirmasi tidak ditemukan")
    await db.bookings.delete_one({"_id": b["_id"]})
    await audit(u["name"], "booking", "reject", bid, f"Tolak booking online {b['client_name']} · {b['date']} {b['start_time']}")
    return {"ok": True, "client_name": b["client_name"], "client_phone": b.get("client_phone", ""), "date": b["date"], "start_time": b["start_time"]}


@api.post("/bookings/{bid}/payments")
async def add_payment(bid: str, body: PaymentIn, u=Depends(BOOK)):
    b = await db.bookings.find_one({"_id": oid(bid)})
    if not b:
        raise HTTPException(404, "Booking tidak ditemukan")
    remaining = b["price"] - b.get("paid_total", 0)
    if body.amount > remaining + 0.01:
        raise HTTPException(400, "Nominal melebihi sisa tagihan")
    today = local_date()
    pay = {"amount": body.amount, "method": body.method, "type": body.type, "note": body.note,
           "date": today, "created_at": now_iso(), "by": u["name"]}
    paid = b.get("paid_total", 0) + body.amount
    st_set = {"paid_total": paid, "payment_status": pay_status(paid, b["price"])}
    push = {"payments": pay}
    if stage_of(b) == "order_received":
        st_set["stage"] = "payment_received"
        push["stage_history"] = stage_entry("payment_received", u["name"])
    await db.bookings.update_one({"_id": oid(bid)}, {"$push": push, "$set": st_set})
    label = "DP" if body.type == "dp" else "Pelunasan"
    await db.transactions.insert_one({"type": "income", "category": "booking", "amount": body.amount,
                                      "method": body.method, "description": f"{label} - {b['client_name']} ({b['package_name']})",
                                      "date": today, "booking_id": bid, "payment_type": body.type, "created_by": u["name"], "created_at": now_iso()})
    await audit(u["name"], "payment", "create", bid,
                f"{label} {fmt_rp(body.amount)} ({body.method}) · {b['client_name']} · {b['package_name']}")
    return ser(await db.bookings.find_one({"_id": oid(bid)}))


# ---------- POS transactions ----------
class TxIn(BaseModel):
    type: Literal["income", "expense"]
    category: Literal["penjualan", "booking", "lainnya", "operasional", "sewa"]
    amount: float = Field(gt=0)
    method: Literal["cash", "transfer", "card", "qris"] = "cash"
    description: str = ""
    date: str
    receipt_path: Optional[str] = None


@api.get("/transactions")
async def list_tx(start: str, end: str, _=Depends(REP)):
    await ensure_rent()
    docs = await db.transactions.find({"date": {"$gte": start, "$lte": end}}).sort([("date", -1), ("created_at", -1)]).to_list(2000)
    return [ser(d) for d in docs]


@api.post("/transactions")
async def create_tx(body: TxIn, u=Depends(current_user)):
    if not eff_perms(u)["expense" if body.type == "expense" else "booking"]:
        raise HTTPException(403, "Akses ditolak")
    doc = body.model_dump() | {"created_by": u["name"], "created_at": now_iso()}
    r = await db.transactions.insert_one(doc)
    doc["_id"] = r.inserted_id
    return ser(doc)


@api.delete("/transactions/{tid}")
async def delete_tx(tid: str, u=Depends(SUPER)):
    old = await db.transactions.find_one({"_id": oid(tid)})
    await db.transactions.delete_one({"_id": oid(tid)})
    if old and old.get("booking_id"):
        await audit(u["name"], "payment", "delete", old["booking_id"], f"Hapus transaksi {old.get('description', '')} {fmt_rp(old['amount'])}")
    return {"ok": True}


async def build_report(period: str, date: str):
    prefix = {"daily": date[:10], "monthly": date[:7], "yearly": date[:4]}[period]
    docs = await db.transactions.find({"date": {"$regex": f"^{prefix}"}}).to_list(10000)
    rev = sum(d["amount"] for d in docs if d["type"] == "income")
    ops = sum(d["amount"] for d in docs if d["type"] == "expense" and d["category"] != "sewa")
    rent = sum(d["amount"] for d in docs if d["type"] == "expense" and d["category"] == "sewa")
    by_cat = {}
    for d in docs:
        by_cat[d["category"]] = by_cat.get(d["category"], 0) + d["amount"]
    bdocs = await db.bookings.find({"date": {"$regex": f"^{prefix}"}, "status": {"$nin": ["cancelled", "pending"]}}).to_list(10000)
    other_inc = sum(d["amount"] for d in docs if d["type"] == "income" and d["category"] != "booking")
    turnover = sum(b["price"] for b in bdocs) + other_inc
    dp = sum(d["amount"] for d in docs if d.get("payment_type") == "dp")
    pel = sum(d["amount"] for d in docs if d.get("payment_type") == "pelunasan")
    return {"period": period, "key": prefix, "turnover": turnover, "income": rev, "revenue": rev,
            "operational": ops, "rental": rent, "expenses": ops + rent, "net_income": rev - ops - rent,
            "net_profit": rev - ops - rent, "dp_total": dp, "pelunasan_total": pel, "other_income": other_inc,
            "by_category": by_cat, "transaction_count": len(docs), "booking_count": len(bdocs),
            "transactions": sorted([ser(d) for d in docs], key=lambda x: (x["date"], x.get("created_at", "")))}


@api.get("/reports")
async def report(period: Literal["daily", "monthly", "yearly"], date: str, _=Depends(REP)):
    await ensure_rent()
    r = await build_report(period, date)
    r.pop("transactions")
    return r


# ---------- attendance ----------
def local_date():
    return (datetime.now(timezone.utc) + LOCAL_TZ).strftime("%Y-%m-%d")


@api.get("/attendance/today")
async def att_today(u=Depends(current_user)):
    return ser(await db.attendance.find_one({"user_id": str(u["_id"]), "date": local_date()}))


class LocIn(BaseModel):
    lat: Optional[float] = None
    lng: Optional[float] = None
    accuracy: Optional[float] = None


def haversine_m(lat1, lng1, lat2, lng2) -> float:
    import math
    r = 6371000
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp, dl = math.radians(lat2 - lat1), math.radians(lng2 - lng1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(a))


async def check_geofence(body: Optional[LocIn]) -> Optional[float]:
    st = await get_settings()
    if not st["geofence_enabled"]:
        return None
    if not body or body.lat is None or body.lng is None:
        raise HTTPException(400, "Lokasi diperlukan untuk absensi. Aktifkan GPS lalu coba lagi.")
    dist = haversine_m(body.lat, body.lng, st["studio_lat"], st["studio_lng"])
    radius = st["attendance_radius_m"]
    if dist > radius:
        raise HTTPException(403, f"Anda berada ±{round(dist)} m dari studio. Absensi hanya bisa dilakukan dalam radius {radius} m.")
    return round(dist)


@api.post("/attendance/clock-in")
async def clock_in(body: Optional[LocIn] = None, u=Depends(current_user)):
    d = local_date()
    if await db.attendance.find_one({"user_id": str(u["_id"]), "date": d}):
        raise HTTPException(400, "Sudah absen masuk hari ini")
    dist = await check_geofence(body)
    doc = {"user_id": str(u["_id"]), "name": u["name"], "sub_role": u["sub_role"], "date": d,
           "clock_in": now_iso(), "clock_out": None, "clock_in_distance_m": dist}
    r = await db.attendance.insert_one(doc)
    doc["_id"] = r.inserted_id
    return ser(doc)


@api.post("/attendance/clock-out")
async def clock_out(body: Optional[LocIn] = None, u=Depends(current_user)):
    rec = await db.attendance.find_one({"user_id": str(u["_id"]), "date": local_date()})
    if not rec:
        raise HTTPException(400, "Belum absen masuk")
    if rec.get("clock_out"):
        raise HTTPException(400, "Sudah absen pulang")
    dist = await check_geofence(body)
    await db.attendance.update_one({"_id": rec["_id"]}, {"$set": {"clock_out": now_iso(), "clock_out_distance_m": dist}})
    return ser(await db.attendance.find_one({"_id": rec["_id"]}))


@api.get("/attendance")
async def att_list(date: Optional[str] = None, u=Depends(current_user)):
    q = {"date": date or local_date()}
    if u["role"] == "staff":
        q["user_id"] = str(u["_id"])
    return [ser(d) for d in await db.attendance.find(q).sort("clock_in", 1).to_list(500)]


# ---------- attendance monthly recap ----------
def local_min(iso: Optional[str]) -> Optional[int]:
    if not iso:
        return None
    d = datetime.fromisoformat(iso) + LOCAL_TZ
    return d.hour * 60 + d.minute


async def attendance_recap(month: str):
    st = await get_settings()
    ws, we = to_min(st["work_start"]), to_min(st["work_end"])
    tol, ot_min, rate = st["late_tolerance_min"], st["overtime_min_minutes"], st["overtime_rate_per_hour"]
    users = await db.users.find({"role": {"$ne": "super_admin"}}).sort("name", 1).to_list(500)
    recs = await db.attendance.find({"date": {"$regex": f"^{month}"}}).to_list(20000)
    out = []
    for u in users:
        mine = [r for r in recs if r["user_id"] == str(u["_id"])]
        late_count = late_minutes = ot_total = 0
        for r in mine:
            cin, cout = local_min(r.get("clock_in")), local_min(r.get("clock_out"))
            if cin is not None and cin - ws > tol:
                late_count += 1
                late_minutes += cin - ws
            if cout is not None and cout - we >= ot_min:
                ot_total += cout - we
        out.append({"user_id": str(u["_id"]), "name": u["name"], "sub_role": u["sub_role"], "phone": u.get("phone", ""), "days_worked": len(mine),
                    "late_count": late_count, "late_minutes": late_minutes, "overtime_minutes": ot_total,
                    "overtime_pay": round(ot_total / 60 * rate)})
    return out


@api.get("/attendance/monthly")
async def att_monthly(month: str, _=Depends(ADMIN)):
    return await attendance_recap(month[:7])


# ---------- payroll (super admin) ----------
class PayrollIn(BaseModel):
    base_salary: float = Field(ge=0, default=0)
    turnover_bonus_pct: float = Field(ge=0, le=100, default=0)
    misc_bonus: float = Field(ge=0, default=0)
    overtime_pay_override: Optional[float] = None
    days_worked_override: Optional[int] = Field(ge=0, le=31, default=None)
    meal_allowance_override: Optional[float] = Field(ge=0, default=None)
    kasbon: float = Field(ge=0, default=0)
    denda: float = Field(ge=0, default=0)
    note: str = ""


@api.get("/payroll")
async def payroll(month: str, _=Depends(SUPER)):
    month = month[:7]
    st = await get_settings()
    recap = await attendance_recap(month)
    turnover = (await build_report("monthly", month + "-01"))["turnover"]
    entries = {e["user_id"]: e for e in await db.payroll.find({"month": month}).to_list(1000)}
    rows = []
    for r in recap:
        e = entries.get(r["user_id"])
        if not e:
            prev = await db.payroll.find({"user_id": r["user_id"], "month": {"$lt": month}}).sort("month", -1).to_list(1)
            e = {"base_salary": prev[0]["base_salary"] if prev else 0, "turnover_bonus_pct": prev[0].get("turnover_bonus_pct", 0) if prev else 0,
                 "misc_bonus": 0, "overtime_pay_override": None, "note": "", "saved": False}
        else:
            e["saved"] = True
        days_ov, meal_ov = e.get("days_worked_override"), e.get("meal_allowance_override")
        days = days_ov if days_ov is not None else r["days_worked"]
        meal_auto = days * st["meal_allowance_per_day"]
        meal = meal_ov if meal_ov is not None else meal_auto
        ot = e["overtime_pay_override"] if e.get("overtime_pay_override") is not None else r["overtime_pay"]
        tb = round(turnover * e["turnover_bonus_pct"] / 100)
        kasbon, denda = e.get("kasbon", 0) or 0, e.get("denda", 0) or 0
        gross = e["base_salary"] + meal + ot + tb + e["misc_bonus"]
        rows.append({**r, "days_worked_auto": r["days_worked"], "days_worked": days, "days_worked_override": days_ov,
                     "base_salary": e["base_salary"], "meal_allowance_auto": meal_auto, "meal_allowance": meal,
                     "meal_allowance_override": meal_ov, "overtime_pay_auto": r["overtime_pay"],
                     "overtime_pay": ot, "overtime_pay_override": e.get("overtime_pay_override"),
                     "turnover_bonus_pct": e["turnover_bonus_pct"], "turnover_bonus": tb, "misc_bonus": e["misc_bonus"],
                     "kasbon": kasbon, "denda": denda, "deductions": kasbon + denda, "gross": gross,
                     "note": e.get("note", ""), "saved": e["saved"], "total": gross - kasbon - denda})
    return {"month": month, "turnover": turnover, "meal_allowance_per_day": st["meal_allowance_per_day"],
            "rows": rows, "grand_total": sum(x["total"] for x in rows)}


@api.put("/payroll/{uid}")
async def save_payroll(uid: str, month: str, body: PayrollIn, _=Depends(SUPER)):
    await db.payroll.update_one({"user_id": uid, "month": month[:7]},
                                {"$set": {**body.model_dump(), "updated_at": now_iso()}}, upsert=True)
    return {"ok": True}


# ---------- object storage (receipts) ----------
STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
APP_NAME = "casadella-studio"
storage_key = None


def init_storage():
    global storage_key
    if storage_key:
        return storage_key
    resp = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": os.environ["EMERGENT_LLM_KEY"]}, timeout=30)
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key


def put_object(path: str, data: bytes, content_type: str) -> dict:
    global storage_key
    for attempt in range(2):
        key = init_storage()
        resp = requests.put(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key, "Content-Type": content_type},
                            data=data, timeout=120)
        if resp.status_code == 503 and attempt == 0:
            storage_key = None
            continue
        resp.raise_for_status()
        return resp.json()


def get_object(path: str):
    key = init_storage()
    resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    resp.raise_for_status()
    return resp.content, resp.headers.get("Content-Type", "application/octet-stream")


@api.post("/upload")
async def upload(file: UploadFile = File(...), u=Depends(EXP)):
    data = await file.read()
    if len(data) > 8 * 1024 * 1024:
        raise HTTPException(400, "Ukuran file maksimal 8MB")
    ext = (file.filename or "img.jpg").rsplit(".", 1)[-1].lower()[:5] or "jpg"
    path = f"{APP_NAME}/uploads/{u['_id']}/{uuid.uuid4()}.{ext}"
    try:
        res = await run_in_threadpool(put_object, path, data, file.content_type or "image/jpeg")
    except requests.HTTPError as e:
        code = e.response.status_code if e.response is not None else 500
        raise HTTPException(402 if code == 402 else 502, "Penyimpanan penuh" if code == 402 else "Gagal mengunggah foto")
    await db.files.insert_one({"path": res["path"], "owner_id": str(u["_id"]), "content_type": file.content_type,
                               "original_name": file.filename, "created_at": now_iso()})
    return {"path": res["path"]}


@api.get("/files/{path:path}")
async def get_file(path: str, _=Depends(REP)):
    f = await db.files.find_one({"path": path})
    if not f:
        raise HTTPException(404, "File tidak ditemukan")
    data, ct = await run_in_threadpool(get_object, path)
    return Response(content=data, media_type=f.get("content_type") or ct, headers={"Cache-Control": "private, max-age=86400"})


# ---------- report export ----------
PERIOD_LABEL = {"daily": "Harian", "monthly": "Bulanan", "yearly": "Tahunan"}
CAT_LABEL = {"booking": "Booking", "penjualan": "Penjualan", "lainnya": "Lainnya", "operasional": "Operasional", "sewa": "Sewa"}


def fmt_rp(n):
    return "Rp " + f"{round(n or 0):,}".replace(",", ".")


def summary_rows(r):
    return [("Total Omzet", r["turnover"]), ("Total Pemasukan", r["income"]), ("  - DP", r["dp_total"]),
            ("  - Pelunasan", r["pelunasan_total"]), ("Biaya Operasional", r["operational"]), ("Biaya Sewa", r["rental"]),
            ("Total Pengeluaran", r["expenses"]), ("Laba Bersih", r["net_profit"])]


def tx_label(t):
    if t.get("payment_type"):
        return "DP" if t["payment_type"] == "dp" else "Pelunasan"
    return CAT_LABEL.get(t["category"], t["category"])


def build_xlsx(r, title, recap):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    for k, v in summary_rows(r):
        ws.append([k, v])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 18
    t = wb.create_sheet("Transaksi")
    t.append(["Tanggal", "Jenis", "Kategori", "Keterangan", "Metode", "Nominal"])
    for x in r["transactions"]:
        t.append([x["date"], "Pemasukan" if x["type"] == "income" else "Pengeluaran", tx_label(x), x.get("description", ""),
                  x.get("method", ""), x["amount"] if x["type"] == "income" else -x["amount"]])
    for col, w in zip("ABCDEF", [12, 12, 14, 40, 10, 14]):
        t.column_dimensions[col].width = w
    if recap is not None:
        a = wb.create_sheet("Absensi")
        a.append(["Nama", "Hari Kerja", "Terlambat (x)", "Terlambat (menit)", "Lembur (menit)", "Uang Lembur"])
        for x in recap:
            a.append([x["name"], x["days_worked"], x["late_count"], x["late_minutes"], x["overtime_minutes"], x["overtime_pay"]])
        a.column_dimensions["A"].width = 24
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(r, title, recap):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    st = getSampleStyleSheet()
    navy = colors.HexColor("#0B3D66")
    els = []
    logo = ROOT_DIR / "logo.png"
    if logo.exists():
        els.append(RLImage(str(logo), width=180, height=180 * 0.213))
    els += [Spacer(1, 8), Paragraph(title, st["Title"]), Spacer(1, 8)]

    def table(data, widths, bold_last=False):
        tb = Table(data, colWidths=widths, repeatRows=1)
        style = [("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                 ("FONTSIZE", (0, 0), (-1, -1), 9), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
                 ("ALIGN", (-1, 1), (-1, -1), "RIGHT")]
        if bold_last:
            style.append(("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"))
        tb.setStyle(TableStyle(style))
        return tb
    els.append(table([["Ringkasan", "Nominal"]] + [[k, fmt_rp(v)] for k, v in summary_rows(r)], [300, 200], True))
    els += [Spacer(1, 14), Paragraph("Transaksi", st["Heading2"])]
    rows = [["Tanggal", "Kategori", "Keterangan", "Metode", "Nominal"]]
    for x in r["transactions"]:
        sign = "" if x["type"] == "income" else "-"
        rows.append([x["date"], tx_label(x), Paragraph(x.get("description", "") or "-", st["BodyText"]), x.get("method", ""), sign + fmt_rp(x["amount"])])
    if len(rows) == 1:
        rows.append(["-", "-", "Tidak ada transaksi", "-", "-"])
    els.append(table(rows, [62, 70, 220, 50, 90]))
    if recap is not None:
        els += [Spacer(1, 14), Paragraph("Rekap Absensi", st["Heading2"])]
        ar = [["Nama", "Hari", "Telat (x)", "Telat (mnt)", "Lembur (mnt)", "Uang Lembur"]]
        ar += [[x["name"], x["days_worked"], x["late_count"], x["late_minutes"], x["overtime_minutes"], fmt_rp(x["overtime_pay"])] for x in recap]
        els.append(table(ar, [150, 40, 60, 70, 75, 95]))
    doc.build(els)
    return buf.getvalue()


@api.get("/reports/export")
async def export_report(period: Literal["daily", "monthly", "yearly"], date: str, format: Literal["pdf", "xlsx"], _=Depends(REP)):
    r = await build_report(period, date)
    st = await get_settings()
    title = f"Laporan {PERIOD_LABEL[period]} {r['key']} - {st['studio_name']}"
    recap = await attendance_recap(r["key"]) if period == "monthly" else None
    fn = f"laporan-{period}-{r['key']}.{format}"
    if format == "pdf":
        data, mt = await run_in_threadpool(build_pdf, r, title, recap), "application/pdf"
    else:
        data, mt = await run_in_threadpool(build_xlsx, r, title, recap), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return Response(content=data, media_type=mt, headers={"Content-Disposition": f'attachment; filename="{fn}"'})


@app.on_event("startup")
async def startup_storage():
    try:
        await run_in_threadpool(init_storage)
    except Exception as e:
        logger.error(f"Storage init failed: {e}")

# ---------- generic PDF ----------
def pdf_doc(title: str, blocks: list, subtitle: str = "") -> bytes:
    """blocks: list of (heading|None, rows[header first], col widths, bold_last)"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    st = getSampleStyleSheet()
    navy = colors.HexColor("#0B3D66")
    els = []
    logo = ROOT_DIR / "logo.png"
    if logo.exists():
        els.append(RLImage(str(logo), width=180, height=180 * 0.213))
    els += [Spacer(1, 8), Paragraph(title, st["Title"])]
    if subtitle:
        els.append(Paragraph(subtitle, st["Normal"]))
    for heading, rows, widths, bold_last in blocks:
        els.append(Spacer(1, 12))
        if heading:
            els.append(Paragraph(heading, st["Heading2"]))
        data = [[Paragraph(str(c), st["BodyText"]) if isinstance(c, str) and len(c) > 40 else c for c in r] for r in rows]
        tb = Table(data, colWidths=widths, repeatRows=1)
        style = [("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                 ("FONTSIZE", (0, 0), (-1, -1), 9), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
                 ("ALIGN", (-1, 1), (-1, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
        if bold_last:
            style.append(("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"))
        tb.setStyle(TableStyle(style))
        els.append(tb)
    doc.build(els)
    return buf.getvalue()


def pdf_response(data: bytes, fn: str):
    return Response(content=data, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{fn}"'})


# ---------- clients (derived from bookings) ----------
async def client_viewer(u=Depends(current_user)):
    p = eff_perms(u)
    if not (p["booking"] or p["reports"]):
        raise HTTPException(403, "Akses ditolak")
    return u


def norm_phone(p: str) -> str:
    p = re.sub(r"\D", "", p or "")
    if p.startswith("0"):
        return "62" + p[1:]
    return "62" + p if p.startswith("8") else p


def client_key(b) -> str:
    return norm_phone(b.get("client_phone")) or "n-" + re.sub(r"\s+", "-", b["client_name"].strip().lower())


async def build_clients():
    bs = await db.bookings.find().sort([("date", 1), ("start_time", 1)]).to_list(50000)
    t = local_date()
    m = {}
    for b in bs:
        k = client_key(b)
        c = m.setdefault(k, {"key": k, "name": b["client_name"], "phone": "", "total_bookings": 0, "total_spent": 0,
                             "total_paid": 0, "last_visit": None, "last_package": None, "next_booking": None})
        c["name"] = b["client_name"]
        if b.get("client_phone"):
            c["phone"] = b["client_phone"]
        c["total_bookings"] += 1
        c["total_spent"] += b["price"]
        c["total_paid"] += b.get("paid_total", 0)
        if b["date"] <= t:
            c["last_visit"], c["last_package"] = b["date"], b["package_name"]
        elif not c["next_booking"]:
            c["next_booking"] = b["date"]
    for p in await db.clients.find({}, {"_id": 0}).to_list(50000):
        c = m.setdefault(p["key"], {"key": p["key"], "total_bookings": 0, "total_spent": 0, "total_paid": 0,
                                    "last_visit": None, "last_package": None, "next_booking": None})
        c.update({f: p.get(f, "") for f in CLIENT_FIELDS})
        c["manual"] = True
    return sorted(m.values(), key=lambda c: c["last_visit"] or c["next_booking"] or c.get("created_at") or "", reverse=True), bs


CLIENT_FIELDS = ["name", "phone", "instagram", "address", "notes", "created_at"]


class ClientIn(BaseModel):
    name: str = Field(min_length=2, max_length=80)
    phone: str = Field(default="", max_length=20)
    instagram: str = Field(default="", max_length=60)
    address: str = Field(default="", max_length=200)
    notes: str = Field(default="", max_length=500)


async def client_exists(key: str) -> bool:
    out, _bs = await build_clients()
    return any(c["key"] == key for c in out)


@api.post("/clients")
async def create_client(body: ClientIn, u=Depends(ADMIN)):
    data = body.model_dump() | {"name": body.name.strip(), "phone": body.phone.strip()}
    key = client_key({"client_phone": data["phone"], "client_name": data["name"]})
    if await client_exists(key):
        raise HTTPException(409, "Klien dengan nomor/nama ini sudah ada")
    await db.clients.insert_one(data | {"key": key, "created_at": now_iso(), "created_by": u["name"]})
    await audit(u["name"], "client", "create", key, f"Klien baru {data['name']} ({data['phone'] or 'tanpa nomor'})")
    return await client_detail(key)


@api.put("/clients/{key}")
async def update_client(key: str, body: ClientIn, u=Depends(ADMIN)):
    out, bs = await build_clients()
    old = next((c for c in out if c["key"] == key), None)
    if not old:
        raise HTTPException(404, "Klien tidak ditemukan")
    data = body.model_dump() | {"name": body.name.strip(), "phone": body.phone.strip()}
    new_key = client_key({"client_phone": data["phone"], "client_name": data["name"]})
    if new_key != key and any(c["key"] == new_key for c in out):
        raise HTTPException(409, "Nomor/nama ini sudah dipakai klien lain")
    ids = [b["_id"] for b in bs if client_key(b) == key]
    if ids:
        await db.bookings.update_many({"_id": {"$in": ids}}, {"$set": {"client_name": data["name"], "client_phone": data["phone"]}})
    prev = await db.clients.find_one_and_delete({"key": key}) or {}
    await db.clients.insert_one(data | {"key": new_key, "created_at": prev.get("created_at") or now_iso(),
                                        "created_by": prev.get("created_by") or u["name"], "updated_by": u["name"], "updated_at": now_iso()})
    labels = {"name": "Nama", "phone": "No. WhatsApp", "instagram": "Instagram", "address": "Alamat", "notes": "Catatan"}
    ch = [{"field": lbl, "from": old.get(f) or "", "to": data[f]} for f, lbl in labels.items() if (old.get(f) or "") != data[f]]
    await audit(u["name"], "client", "update", new_key, f"Ubah data klien {data['name']}", ch)
    return await client_detail(new_key)


@api.get("/clients")
async def list_clients(q: Optional[str] = None, _=Depends(client_viewer)):
    out, _bs = await build_clients()
    if q:
        ql = q.lower()
        out = [c for c in out if ql in c["name"].lower() or ql in (c["phone"] or "")]
    return out


@api.get("/clients/export")
async def export_clients(format: Literal["pdf", "xlsx"], _=Depends(client_viewer)):
    out, _bs = await build_clients()
    head = ["Nama", "No. WhatsApp", "Kunjungan Terakhir", "Paket Terakhir", "Jml Booking", "Total Belanja"]
    if format == "xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Klien"
        ws.append(head)
        for c in out:
            ws.append([c["name"], c["phone"], c["last_visit"] or "-", c["last_package"] or "-", c["total_bookings"], c["total_spent"]])
        for col, w in zip("ABCDEF", [26, 18, 18, 24, 12, 16]):
            ws.column_dimensions[col].width = w
        buf = io.BytesIO()
        wb.save(buf)
        return Response(content=buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": 'attachment; filename="database-klien.xlsx"'})
    rows = [head] + [[c["name"], c["phone"] or "-", c["last_visit"] or "-", c["last_package"] or "-", c["total_bookings"], fmt_rp(c["total_spent"])] for c in out]
    st = await get_settings()
    data = await run_in_threadpool(pdf_doc, f"Database Klien - {st['studio_name']}", [(None, rows, [110, 85, 75, 110, 45, 95], False)], f"Total {len(out)} klien")
    return pdf_response(data, "database-klien.pdf")


@api.get("/clients/{key}")
async def client_detail(key: str, _=Depends(client_viewer)):
    out, bs = await build_clients()
    c = next((x for x in out if x["key"] == key), None)
    if not c:
        raise HTTPException(404, "Klien tidak ditemukan")
    hist = [ser(b) for b in reversed(bs) if client_key(b) == key]
    fav = {}
    for b in hist:
        fav[b["package_name"]] = fav.get(b["package_name"], 0) + 1
    return {**c, "favorite_package": max(fav, key=fav.get) if fav else None, "bookings": hist}


# ---------- package popularity ----------
@api.get("/reports/packages")
async def package_popularity(period: Literal["daily", "monthly", "yearly"], date: str, _=Depends(REP)):
    prefix = {"daily": date[:10], "monthly": date[:7], "yearly": date[:4]}[period]
    bs = await db.bookings.find({"date": {"$regex": f"^{prefix}"}}).to_list(20000)
    agg = {}
    for b in bs:
        a = agg.setdefault(b["package_name"], {"package_name": b["package_name"], "count": 0, "revenue": 0})
        a["count"] += 1
        a["revenue"] += b["price"]
    total = sum(a["count"] for a in agg.values())
    out = sorted(agg.values(), key=lambda a: (-a["count"], -a["revenue"]))
    for a in out:
        a["share"] = round(a["count"] / total * 100) if total else 0
    return out


# ---------- invoice & payslip PDFs ----------
@api.get("/bookings/{bid}/invoice")
async def booking_invoice(bid: str, _=Depends(client_viewer)):
    b = await db.bookings.find_one({"_id": oid(bid)})
    if not b:
        raise HTTPException(404, "Booking tidak ditemukan")
    st = await get_settings()
    inv = f"INV-{b['date'].replace('-', '')}-{str(b['_id'])[-5:].upper()}"
    info = [["Detail", "Keterangan"], ["Klien", b["client_name"]], ["No. WhatsApp", b.get("client_phone") or "-"],
            ["Tanggal Sesi", b["date"]], ["Jam", f"{b['start_time']} - {b['end_time']}"], ["Fotografer", b.get("photographer_name") or "-"]]
    items = [["Paket", "Durasi", "Harga"], [b["package_name"], f"{b['duration_min']} menit", fmt_rp(b["price"])]]
    pays = [["Tanggal", "Jenis", "Metode", "Nominal"]] + [[p["date"], "DP" if p["type"] == "dp" else "Pelunasan", p["method"], fmt_rp(p["amount"])] for p in b.get("payments", [])]
    if len(pays) == 1:
        pays.append(["-", "-", "-", fmt_rp(0)])
    status = {"lunas": "LUNAS", "dp": "DP", "belum_bayar": "BELUM BAYAR"}[b.get("payment_status", "belum_bayar")]
    dp_total = sum(p["amount"] for p in b.get("payments", []) if p["type"] == "dp")
    summ = [["Ringkasan", "Nominal"], ["Total", fmt_rp(b["price"])], ["DP", fmt_rp(dp_total)], ["Sudah Dibayar", fmt_rp(b.get("paid_total", 0))],
            [f"Sisa Tagihan ({status})", fmt_rp(b["price"] - b.get("paid_total", 0))]]
    data = await run_in_threadpool(pdf_doc, f"Invoice {inv}", [(None, info, [150, 350], False), ("Item", items, [250, 100, 150], False),
                                                               ("Pembayaran", pays, [110, 110, 110, 170], False), (None, summ, [300, 200], True)],
                                   st["studio_name"])
    return pdf_response(data, f"{inv}.pdf")


@api.get("/payroll/{uid}/payslip")
async def payslip(uid: str, month: str, u=Depends(SUPER)):
    p = await payroll(month, u)
    r = next((x for x in p["rows"] if x["user_id"] == uid), None)
    if not r:
        raise HTTPException(404, "Karyawan tidak ditemukan")
    st = await get_settings()
    rows = [["Komponen", "Nominal"], ["Gaji Pokok", fmt_rp(r["base_salary"])],
            [f"Uang Makan ({r['days_worked']} hari)", fmt_rp(r["meal_allowance"])],
            [f"Uang Lembur ({r['overtime_minutes']} menit)", fmt_rp(r["overtime_pay"])],
            [f"Bonus Omzet ({r['turnover_bonus_pct']}%)", fmt_rp(r["turnover_bonus"])],
            ["Bonus Lain-lain", fmt_rp(r["misc_bonus"])], ["Total Pendapatan", fmt_rp(r["gross"])],
            ["Potongan Kasbon", "-" + fmt_rp(r["kasbon"])], ["Potongan Denda", "-" + fmt_rp(r["denda"])],
            ["TOTAL DITERIMA", fmt_rp(r["total"])]]
    att = [["Kehadiran", "Jumlah"], ["Hari Kerja", r["days_worked"]], ["Terlambat", f"{r['late_count']}x ({r['late_minutes']} menit)"],
           ["Lembur", f"{r['overtime_minutes']} menit"]]
    blocks = [(None, rows, [330, 170], True), ("Rekap Absensi", att, [330, 170], False)]
    if r.get("note"):
        blocks.append((None, [["Catatan"], [r["note"]]], [500], False))
    data = await run_in_threadpool(pdf_doc, f"Slip Gaji {p['month']}", blocks, f"{r['name']} - {st['studio_name']}")
    return pdf_response(data, f"slip-gaji-{p['month']}-{re.sub(r'[^a-z0-9]+', '-', r['name'].lower())}.pdf")


# ---------- audit log (super admin) ----------
@api.get("/audit")
async def list_audit(entity: Optional[Literal["booking", "payment", "package", "client"]] = None, limit: int = Query(200, le=1000), _=Depends(SUPER)):
    q = {"entity": entity} if entity else {}
    return [ser(d) for d in await db.audit_logs.find(q).sort("created_at", -1).to_list(limit)]


# ---------- monthly rent (auto) ----------
async def ensure_rent():
    st = await get_settings()
    if not st["rent_auto_enabled"] or not st["rent_amount"]:
        return
    month = local_date()[:7]
    r = await db.rent_runs.update_one({"month": month}, {"$setOnInsert": {"month": month, "created_at": now_iso()}}, upsert=True)
    if r.upserted_id:
        await db.transactions.insert_one({"type": "expense", "category": "sewa", "amount": st["rent_amount"], "method": "transfer",
                                          "description": "Sewa tempat bulanan (otomatis)", "date": f"{month}-01", "auto_rent": month,
                                          "created_by": "Sistem", "created_at": now_iso()})


async def rent_loop():
    while True:
        try:
            await ensure_rent()
        except Exception as e:
            logger.error(f"Auto rent failed: {e}")
        await asyncio.sleep(3600)


@app.on_event("startup")
async def start_rent_loop():
    asyncio.create_task(rent_loop())


# ---------- public online booking ----------
class PublicBookingIn(BaseModel):
    client_name: str = Field(min_length=2, max_length=80)
    client_phone: str = Field(min_length=8, max_length=20)
    date: str = Field(pattern=r"^\d{4}-\d{2}-\d{2}$")
    start_time: str = Field(pattern=r"^\d{2}:\d{2}$")
    package_id: str
    notes: str = Field(default="", max_length=300)


def now_local_min():
    n = datetime.now(timezone.utc) + LOCAL_TZ
    return n.hour * 60 + n.minute


async def public_package(pid: str):
    p = await db.packages.find_one({"_id": oid(pid), "active": True})
    if not p:
        raise HTTPException(404, "Paket tidak ditemukan")
    return p


async def public_slots_for(date: str, p):
    if date < local_date():
        return []
    sl = await compute_slots(date, p["duration_min"])
    if date == local_date():
        cut = now_local_min() + 30
        for s in sl:
            if to_min(s["time"]) < cut:
                s["available"] = False
    return sl


@api.get("/public/info")
async def public_info():
    st = await get_settings()
    pk = await db.packages.find({"active": True}).sort("price", 1).to_list(200)
    return {"studio_name": st["studio_name"], "open_time": st["open_time"], "close_time": st["close_time"], "today": local_date(),
            "packages": [{"id": str(p["_id"]), "name": p["name"], "description": p.get("description", ""),
                          "duration_min": p["duration_min"], "price": p["price"]} for p in pk]}


@api.get("/public/slots")
async def public_slots(date: str, package_id: str):
    return await public_slots_for(date, await public_package(package_id))


@api.post("/public/bookings")
async def public_create_booking(body: PublicBookingIn):
    p = await public_package(body.package_id)
    sl = await public_slots_for(body.date, p)
    match = next((x for x in sl if x["time"] == body.start_time), None)
    if not match or not match["available"]:
        raise HTTPException(409, "Slot waktu sudah tidak tersedia, silakan pilih jam lain")
    doc = await build_booking(BookingIn(client_name=body.client_name, client_phone=body.client_phone, date=body.date,
                                        start_time=body.start_time, package_id=body.package_id, notes=body.notes))
    doc.update({"payments": [], "paid_total": 0, "payment_status": "belum_bayar", "status": "pending", "source": "online",
                "created_by": "Klien (online)", "created_at": now_iso()})
    r = await db.bookings.insert_one(doc)
    await audit("Klien (online)", "booking", "create", str(r.inserted_id),
                f"Booking online {doc['client_name']} · {doc['date']} {doc['start_time']} · {doc['package_name']}")
    return {"id": str(r.inserted_id), "client_name": doc["client_name"], "date": doc["date"], "start_time": doc["start_time"],
            "end_time": doc["end_time"], "package_name": doc["package_name"], "price": doc["price"], "status": "pending"}


@api.get("/public/logo")
async def public_logo():
    return FileResponse(ROOT_DIR / "logo.png", media_type="image/png")


@api.get("/book", response_class=HTMLResponse)
async def booking_page():
    return HTMLResponse((ROOT_DIR / "booking_page.html").read_text(encoding="utf-8"))


# ---------- workflow ----------
STAGES = ["order_received", "payment_received", "check_schedule", "client_arrived", "photo_session",
          "editing", "awaiting_final", "addon", "completed"]
STAGE_LABEL = {"order_received": "Order Received!", "payment_received": "Payment Received!", "check_schedule": "Check Schedule",
               "client_arrived": "Klien Tiba di Studio", "photo_session": "Photo Session", "editing": "Editing in Progress",
               "awaiting_final": "Menunggu Pelunasan", "addon": "Add-on Order Received!", "completed": "{nama} - Paid in Full"}
STAGE_TEXT = dict(STAGE_LABEL)  # live labels (editable by super admin), loaded at startup


def label_of(stage: str, client: str = "") -> str:
    return STAGE_TEXT[stage].replace("{nama}", client)


def stage_of(b) -> str:
    return b.get("stage") or ("payment_received" if b.get("paid_total", 0) > 0 else "order_received")


def stage_idx(b) -> int:
    return STAGES.index(stage_of(b))


def add_working_days(d: str, n: int) -> str:
    cur = datetime.strptime(d, "%Y-%m-%d")
    while n > 0:
        cur += timedelta(days=1)
        if cur.weekday() != 6:  # Sunday off
            n -= 1
    return cur.strftime("%Y-%m-%d")


def stage_entry(stage: str, by: str):
    return {"stage": stage, "at": now_iso(), "by": by}


def wf_view(b) -> dict:
    d = ser(b)
    st = stage_of(b)
    d["stage"] = st
    d["stage_label"] = label_of(st, b["client_name"])
    t = local_date()
    d["overdue"] = (b["date"] < t and STAGES.index(st) < STAGES.index("editing")) or \
                   (st == "editing" and (b.get("edit_deadline") or "9999") < t)
    return d


class StageIn(BaseModel):
    action: Literal["check_schedule", "client_arrived", "start_session", "finish_session", "finish_editing", "complete_order"]


ACTIONS = {  # action: (target stage, allowed current stages, who)
    "check_schedule": ("check_schedule", {"order_received", "payment_received"}, "admin"),
    "client_arrived": ("client_arrived", {"order_received", "payment_received", "check_schedule"}, "admin"),
    "start_session": ("photo_session", {"client_arrived"}, "photographer"),
    "finish_session": ("editing", {"photo_session"}, "photographer"),
    "finish_editing": ("awaiting_final", {"editing"}, "editor"),
    "complete_order": ("completed", {"awaiting_final", "addon"}, "admin"),
}
ACTION_HINT = {"start_session": "Menunggu admin menandai Klien Datang", "finish_session": "Mulai sesi foto terlebih dahulu",
               "finish_editing": "Belum ada pekerjaan edit untuk booking ini", "complete_order": "Order belum sampai tahap pelunasan"}


def may_act(u, b, who: str) -> bool:
    if eff_perms(u)["booking"]:
        return True
    if who == "photographer":
        return u["sub_role"] == "photographer" and b.get("photographer_id") in (None, "", str(u["_id"]))
    if who == "editor":
        return u["sub_role"] == "editor" and b.get("editor_id") in (None, "", str(u["_id"]))
    return False


@api.post("/bookings/{bid}/stage")
async def advance_stage(bid: str, body: StageIn, u=Depends(current_user)):
    b = await db.bookings.find_one({"_id": oid(bid)})
    if not b or b.get("status") != "confirmed":
        raise HTTPException(404, "Booking aktif tidak ditemukan")
    target, allowed, who = ACTIONS[body.action]
    if not may_act(u, b, who):
        raise HTTPException(403, "Akses ditolak")
    cur = stage_of(b)
    if cur not in allowed:
        raise HTTPException(400, ACTION_HINT.get(body.action, f"Tidak bisa dari tahap {label_of(cur, b['client_name'])}"))
    upd = {"stage": target}
    if body.action == "complete_order" and b.get("paid_total", 0) < b["price"]:
        raise HTTPException(400, "Order belum lunas. Catat pelunasan terlebih dahulu.")
    if body.action == "finish_session":
        received = local_date()
        upd.update({"photo_done_by": u["name"], "photo_done_at": now_iso(), "edit_received_at": now_iso(),
                    "edit_received_date": received, "edit_deadline": add_working_days(received, 3)})
    if body.action == "start_session":
        upd.update({"photo_started_by": u["name"], "photo_started_at": now_iso()})
    if body.action == "finish_editing":
        upd.update({"edit_done_by": u["name"], "edit_done_at": now_iso(), "edit_done_date": local_date()})
    await db.bookings.update_one({"_id": b["_id"]}, {"$set": upd, "$push": {"stage_history": stage_entry(target, u["name"])}})
    await audit(u["name"], "booking", "stage", bid, f"{b['client_name']}: {label_of(cur, b['client_name'])} → {label_of(target, b['client_name'])}")
    return wf_view(await db.bookings.find_one({"_id": b["_id"]}))


class AddonIn(BaseModel):
    name: str = Field(min_length=1, max_length=80)
    price: float = Field(gt=0)


@api.post("/bookings/{bid}/addons")
async def add_addon(bid: str, body: AddonIn, u=Depends(BOOK)):
    b = await db.bookings.find_one({"_id": oid(bid)})
    if not b or b.get("status") != "confirmed":
        raise HTTPException(404, "Booking aktif tidak ditemukan")
    if stage_of(b) not in ("awaiting_final", "addon"):
        raise HTTPException(400, "Add-on hanya bisa ditambahkan saat pelunasan (setelah editing selesai)")
    price = b["price"] + body.price
    addon = {"name": body.name.strip(), "price": body.price, "by": u["name"], "at": now_iso()}
    upd = {"price": price, "payment_status": pay_status(b.get("paid_total", 0), price), "stage": "addon"}
    push = {"addons": addon}
    if stage_of(b) != "addon":
        push = {"addons": addon, "stage_history": stage_entry("addon", u["name"])}
    await db.bookings.update_one({"_id": b["_id"]}, {"$set": upd, "$push": push})
    await audit(u["name"], "booking", "addon", bid, f"Add-on {addon['name']} {fmt_rp(body.price)} · {b['client_name']}",
                [{"field": "Harga", "from": b["price"], "to": price}])
    return wf_view(await db.bookings.find_one({"_id": b["_id"]}))


@api.get("/workflow")
async def workflow(include_completed: bool = False, _=Depends(ADMIN)):
    q = {"status": "confirmed"}
    if not include_completed:
        q["stage"] = {"$ne": "completed"}
    docs = await db.bookings.find(q).sort([("date", 1), ("start_time", 1)]).to_list(2000)
    items = [wf_view(b) for b in docs]
    counts = {s: 0 for s in STAGES}
    for i in items:
        counts[i["stage"]] += 1
    return {"stages": [{"key": s, "label": STAGE_TEXT[s], "count": counts[s]} for s in STAGES],
            "items": items, "overdue": sum(1 for i in items if i["overdue"])}


@api.get("/workflow/labels")
async def get_labels(_=Depends(current_user)):
    return [{"key": s, "label": STAGE_TEXT[s], "default": STAGE_LABEL[s]} for s in STAGES]


@api.put("/workflow/labels")
async def put_labels(body: dict, u=Depends(SUPER)):
    new = {k: str(v).strip()[:80] for k, v in body.items() if k in STAGE_LABEL and str(v).strip()}
    ch = [{"field": k, "from": STAGE_TEXT[k], "to": v} for k, v in new.items() if STAGE_TEXT[k] != v]
    STAGE_TEXT.update(new)
    await db.settings.update_one({"key": "stage_labels"}, {"$set": {"labels": dict(STAGE_TEXT)}}, upsert=True)
    if ch:
        await audit(u["name"], "booking", "labels", "workflow", "Ubah teks workflow", ch)
    return await get_labels(u)


@app.on_event("startup")
async def load_labels():
    doc = await db.settings.find_one({"key": "stage_labels"})
    if doc:
        STAGE_TEXT.update({k: v for k, v in doc.get("labels", {}).items() if k in STAGE_LABEL})


# ---------- jobs (photographer / editor) ----------
def job_of(b, kind: str) -> dict:
    v = wf_view(b)
    base = {"booking_id": v["id"], "kind": kind, "client_name": b["client_name"], "date": b["date"], "start_time": b["start_time"],
            "package_name": b["package_name"], "stage": v["stage"], "stage_label": v["stage_label"]}
    if kind == "photo":
        done = stage_idx(b) >= STAGES.index("editing")
        return {**base, "assignee_id": b.get("photographer_id"), "assignee_name": b.get("photographer_name") or "Semua Fotografer",
                "done": done, "done_by": b.get("photo_done_by"), "done_date": (b.get("photo_done_at") or "")[:10],
                "due": b["date"], "overdue": not done and b["date"] < local_date()}
    done = stage_idx(b) >= STAGES.index("awaiting_final")
    return {**base, "assignee_id": b.get("editor_id"), "assignee_name": b.get("editor_name") or "Semua Editor",
            "done": done, "done_by": b.get("edit_done_by"), "done_date": b.get("edit_done_date") or "",
            "received": b.get("edit_received_date"), "due": b.get("edit_deadline"),
            "overdue": not done and (b.get("edit_deadline") or "9999") < local_date()}


async def jobs_for(u, d: str) -> list:
    me = str(u["_id"])
    if u["sub_role"] == "photographer":
        docs = await db.bookings.find({"status": "confirmed", "photographer_id": {"$in": [None, "", me]},
                                       "$or": [{"date": d}, {"date": {"$lt": d}, "stage": {"$nin": STAGES[5:]}}]}).sort("start_time", 1).to_list(500)
        return [job_of(b, "photo") for b in docs]
    if u["sub_role"] == "editor":
        docs = await db.bookings.find({"status": "confirmed", "editor_id": {"$in": [None, "", me]},
                                       "$or": [{"stage": "editing"}, {"edit_done_date": d}]}).sort("edit_deadline", 1).to_list(500)
        return [job_of(b, "edit") for b in docs]
    return []


# ---------- daily checklist ----------
def task_role(sub_role: str) -> Optional[str]:
    return {"admin1": "admin", "admin2": "admin", "photographer": "photographer", "editor": "editor"}.get(sub_role)


class TemplateIn(BaseModel):
    title: str = Field(min_length=2, max_length=100)
    role: Literal["admin", "photographer", "editor"]
    target: int = Field(ge=1, le=100, default=1)


@api.get("/task-templates")
async def list_templates(_=Depends(ADMIN)):
    return [ser(d) for d in await db.task_templates.find({"active": True}).sort([("role", 1), ("order", 1)]).to_list(500)]


@api.post("/task-templates")
async def create_template(body: TemplateIn, u=Depends(SUPER)):
    n = await db.task_templates.count_documents({"role": body.role})
    doc = body.model_dump() | {"title": body.title.strip(), "active": True, "order": n, "created_at": now_iso()}
    r = await db.task_templates.insert_one(doc)
    doc["_id"] = r.inserted_id
    return ser(doc)


@api.delete("/task-templates/{tid}")
async def delete_template(tid: str, _=Depends(SUPER)):
    await db.task_templates.update_one({"_id": oid(tid)}, {"$set": {"active": False}})
    return {"ok": True}


async def tasks_for(uid: str, role: Optional[str], d: str) -> list:
    if not role:
        return []
    tpls = await db.task_templates.find({"role": role, "active": True}).sort("order", 1).to_list(200)
    logs = {x["template_id"]: x for x in await db.task_logs.find({"user_id": uid, "date": d}).to_list(500)}
    out = []
    for t in tpls:
        lg = logs.get(str(t["_id"]), {})
        out.append({"template_id": str(t["_id"]), "title": t["title"], "target": t["target"], "auto": t.get("auto"),
                    "count": lg.get("count", 0), "done": lg.get("done", False), "done_at": lg.get("done_at")})
    return out


@api.get("/checklist/today")
async def checklist_today(u=Depends(current_user)):
    d = local_date()
    tasks = await tasks_for(str(u["_id"]), task_role(u["sub_role"]), d)
    jobs = await jobs_for(u, d)
    req = [t for t in tasks if t["target"] > 0]
    done = sum(1 for t in req if t["done"]) + sum(1 for j in jobs if j["done"])
    return {"date": d, "tasks": tasks, "jobs": jobs, "completed": done, "total": len(req) + len(jobs)}


async def bump_task(u, tid: str, complete: bool, auto: bool = False):
    t = await db.task_templates.find_one({"_id": oid(tid), "active": True})
    if not t or t["role"] != task_role(u["sub_role"]):
        raise HTTPException(404, "Tugas tidak ditemukan")
    if t.get("auto") and not auto:
        raise HTTPException(400, "Tugas ini dihitung otomatis dari booking baru")
    d, uid = local_date(), str(u["_id"])
    lg = await db.task_logs.find_one({"user_id": uid, "date": d, "template_id": tid}) or {}
    if complete and lg.get("done"):
        raise HTTPException(400, "Tugas sudah selesai")
    count = max(t["target"], lg.get("count", 0)) if complete else lg.get("count", 0) + 1
    if complete and t["target"] > 1:
        raise HTTPException(400, f"Tugas ini selesai otomatis setelah {t['target']} kali")
    done = count >= t["target"]
    upd = {"count": count, "done": done, "name": u["name"]}
    if done and not lg.get("done"):
        upd["done_at"] = now_iso()
    elif not done:
        upd["done_at"] = None
    await db.task_logs.update_one({"user_id": uid, "date": d, "template_id": tid}, {"$set": upd}, upsert=True)
    return await checklist_today(u)


@api.post("/checklist/{tid}/increment")
async def checklist_inc(tid: str, u=Depends(current_user)):
    return await bump_task(u, tid, False)


async def count_auto_task(u, kind: str):
    role = task_role(u["sub_role"])
    if not role:
        return
    for t in await db.task_templates.find({"role": role, "active": True, "auto": kind}).to_list(20):
        await bump_task(u, str(t["_id"]), False, auto=True)


@api.post("/checklist/{tid}/complete")
async def checklist_complete(tid: str, u=Depends(current_user)):
    return await bump_task(u, tid, True)


@api.get("/reports/checklist")
async def checklist_report(date: str, _=Depends(ADMIN)):
    users = await db.users.find({"role": {"$ne": "super_admin"}, "enabled": {"$ne": False}}).sort("name", 1).to_list(500)
    out = []
    for u in users:
        tasks = await tasks_for(str(u["_id"]), task_role(u["sub_role"]), date)
        jobs = [j for j in await jobs_for(u, date) if j["done"] and j["done_date"] == date or not j["done"] and (j["due"] or "9999") <= date]
        items = [{"title": f"{t['title']} ({t['count']}/min {t['target']})", "done": t["done"]} for t in tasks if t["target"] > 0]
        items += [{"title": f"{t['title']} ({t['count']} · tanpa batas)", "done": True} for t in tasks if t["target"] == 0 and t["count"]]
        items += [{"title": f"{'Foto' if j['kind'] == 'photo' else 'Edit'}: {j['client_name']} · {j['package_name']}"
                            + (" (lewat deadline)" if j["overdue"] else ""), "done": j["done"]} for j in jobs]
        completed = sum(1 for i in items if i["done"])
        pending = len(items) - completed
        note = ("Great job! Semua tugas selesai." if items and not pending else
                f"Lalai: {pending} tugas belum selesai — indikasi kurang disiplin/malas." if pending else "Tidak ada tugas hari ini.")
        out.append({"user_id": str(u["_id"]), "name": u["name"], "sub_role": u["sub_role"], "completed": completed,
                    "pending": pending, "items": items, "note": note, "status": "great" if items and not pending else "lazy" if pending else "none"})
    return {"date": date, "staff": out, "completed": sum(x["completed"] for x in out), "pending": sum(x["pending"] for x in out)}


# ---------- attendance: late today ----------
@api.get("/attendance/late-today")
async def late_today(_=Depends(ADMIN)):
    st = await get_settings()
    ws, tol, d = to_min(st["work_start"]), st["late_tolerance_min"], local_date()
    users = await db.users.find({"role": {"$ne": "super_admin"}, "enabled": {"$ne": False}}).sort("name", 1).to_list(500)
    recs = {r["user_id"]: r for r in await db.attendance.find({"date": d}).to_list(500)}
    now_m = now_local_min()
    out = []
    for u in users:
        r = recs.get(str(u["_id"]))
        cin = local_min(r.get("clock_in")) if r else None
        if cin is not None and cin - ws > tol:
            out.append({"user_id": str(u["_id"]), "name": u["name"], "sub_role": u["sub_role"], "status": "late",
                        "clock_in": r["clock_in"], "late_minutes": cin - ws})
        elif cin is None and now_m - ws > tol:
            out.append({"user_id": str(u["_id"]), "name": u["name"], "sub_role": u["sub_role"], "status": "absent",
                        "clock_in": None, "late_minutes": now_m - ws})
    return {"work_start": st["work_start"], "items": out}


DEFAULT_TASKS = [("Input booking harian", 0), ("Follow up klien", 10), ("Siapkan konten broadcast", 1), ("Kirim broadcast", 1)]


# ---------- seed ----------
@app.on_event("startup")
async def seed_tasks():
    if await db.task_templates.count_documents({}) == 0:
        await db.task_templates.insert_many([{"title": t, "role": "admin", "target": n, "active": True, "order": i, "created_at": now_iso()}
                                             for i, (t, n) in enumerate(DEFAULT_TASKS)])
    await db.task_templates.update_many({"title": "Input booking harian", "role": "admin", "auto": {"$exists": False}}, {"$set": {"auto": "booking"}})
    await db.task_templates.update_many({"auto": "booking"}, {"$set": {"target": 0}})


@app.on_event("startup")
async def seed():
    if not await db.settings.find_one({"key": "system"}):
        await db.settings.insert_one({"key": "system", "studio_name": "Casadella Photo Studio", "open_time": "09:00",
                                      "close_time": "20:00", "slot_interval": 15, "default_dp_percent": 50})
    if await db.users.count_documents({}) == 0:
        seeds = [("Super Admin", "super_admin", os.environ["SEED_SUPERADMIN_PIN"]),
                 ("Admin 1", "admin1", "111111"), ("Admin 2", "admin2", "222222"),
                 ("Rina (Fotografer)", "photographer", "333333"), ("Budi (Editor)", "editor", "444444")]
        for name, sr, pin in seeds:
            await db.users.insert_one({"name": name, "sub_role": sr, "role": ROLE_MAP[sr], "enabled": True,
                                       "pin_hash": hash_pin(pin), "created_at": now_iso()})
    if await db.packages.count_documents({}) == 0:
        await db.packages.insert_many([
            {"name": "Self Photo Basic", "description": "15 menit, 1-2 orang, 5 foto edit", "duration_min": 15, "price": 75000, "active": True},
            {"name": "Couple Session", "description": "30 menit, 2 orang, 10 foto edit", "duration_min": 30, "price": 250000, "active": True},
            {"name": "Family Portrait", "description": "60 menit, hingga 6 orang, 20 foto edit", "duration_min": 60, "price": 600000, "active": True},
            {"name": "Graduation", "description": "45 menit, toga & properti, 15 foto edit", "duration_min": 45, "price": 400000, "active": True},
            {"name": "Prewedding Studio", "description": "120 menit, 2 konsep, 30 foto edit", "duration_min": 120, "price": 1500000, "active": True},
        ])


app.include_router(api)
app.add_middleware(CORSMiddleware, allow_credentials=True, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])
logging.basicConfig(level=logging.INFO)


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
